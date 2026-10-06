import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Arial'
        run.font.color.rgb = RGBColor(0, 51, 102) # Dark blue similar to the example

def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    if bold:
        run.bold = True
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)

def add_screenshot_placeholder(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("[MASUKKAN SCREENSHOT DI SINI]")
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.italic = True
    
    # Add an empty box shape (using a table with 1 cell as a placeholder box)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = table.cell(0, 0)
    # Just to add some space
    cell.text = "\n\n\n\n\n"
    # Borders can be set using XML if needed, but we'll leave it simple
    doc.add_paragraph() # Add space after

doc = docx.Document()

# --- TITLE PAGE ---
doc.add_paragraph('\n\n\n\n\n')
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run("BUKU PANDUAN PENGGUNAAN\n\nWEBSITE PRANATA")
title_run.font.name = 'Arial'
title_run.font.size = Pt(24)
title_run.bold = True
title_run.font.color.rgb = RGBColor(0, 51, 102)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run = subtitle.add_run("(Sistem Pengelolaan Data, Visualisasi, dan Narasi AI Terintegrasi)")
subtitle_run.font.name = 'Arial'
subtitle_run.font.size = Pt(14)
subtitle_run.font.color.rgb = RGBColor(100, 100, 100)
doc.add_page_break()

# --- DAFTAR ISI ---
add_heading(doc, 'DAFTAR ISI', level=1)
toc_items = [
    "PENDAHULUAN",
    "BAB 1: PERAN DAN HAK AKSES PENGGUNA",
    "BAB 2: MODUL AUTENTIKASI DAN HALAMAN UTAMA",
    "BAB 3: PANDUAN OPERASIONAL ADMIN (ROLE 1)",
    "BAB 4: PANDUAN OPERASIONAL PENANGGUNG JAWAB (ROLE 3)",
    "BAB 5: PANDUAN OPERASIONAL PENGGUNA UMUM (ROLE 2 & 4)"
]
for item in toc_items:
    add_paragraph(doc, item)
doc.add_page_break()

# --- PENDAHULUAN ---
add_heading(doc, 'PENDAHULUAN', level=1)
add_paragraph(doc, 'Buku panduan ini disusun untuk memberikan petunjuk operasional yang terperinci dan terstruktur dalam penggunaan website PRANATA.')
add_paragraph(doc, 'PRANATA dikembangkan untuk mentransformasikan proses pengelolaan data, pembuatan laporan narasi berbasis Artificial Intelligence (AI), dan visualisasi data secara terintegrasi.')
add_paragraph(doc, 'Melalui sistem ini, diharapkan proses manajemen data dan pelaporan menjadi lebih efektif, efisien, transparan, serta dapat menyajikan informasi yang akurat dengan bantuan teknologi AI.')
doc.add_page_break()

# --- BAB 1 ---
add_heading(doc, 'BAB 1: PERAN DAN HAK AKSES PENGGUNA', level=1)
add_paragraph(doc, 'Website PRANATA dirancang dengan membagi hak akses ke dalam 3 (tiga) kelompok aktor utama. Setiap aktor memiliki peran, batasan otoritas, serta fungsi menu yang spesifik untuk menjaga integritas dan keamanan data operasional.')

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'No'
hdr_cells[1].text = 'Aktor (Role)'
hdr_cells[2].text = 'Deskripsi & Tanggung Jawab Utama'

roles_data = [
    ('1', 'Admin\n(Role 1)', 'Pengelola sistem utama. Memiliki akses penuh terhadap manajemen pengguna, kelola data master (kategori, subjek, indikator), manajemen pengetahuan AI, pengaturan model AI, dan akses penuh ke dashboard serta peta.'),
    ('2', 'Penanggung Jawab\n(Role 3)', 'Pengelola konten dan operasional. Memiliki akses untuk mengelola data (kategori, subjek, indikator), manajemen pengetahuan AI, serta melihat dashboard dan menggunakan fitur generate narasi AI. Tidak dapat mengelola pengguna.'),
    ('3', 'Pengguna Umum\n(Role 2 & 4)', 'Pengguna akhir sistem. Hanya dapat melihat dashboard, melakukan pemfilteran data, dan melihat data (view only) tanpa akses untuk mengubah data.')
]

for row_num, role, desc in roles_data:
    row_cells = table.add_row().cells
    row_cells[0].text = row_num
    row_cells[1].text = role
    row_cells[2].text = desc

doc.add_page_break()

# --- BAB 2 ---
add_heading(doc, 'BAB 2: MODUL AUTENTIKASI DAN HALAMAN UTAMA', level=1)
add_heading(doc, '2.1 Login (Autentikasi Pengguna)', level=2)
add_paragraph(doc, '1. Buka peramban web dan arahkan ke alamat URL aplikasi PRANATA.')
add_paragraph(doc, '2. Sistem akan menampilkan halaman utama login.')
add_screenshot_placeholder(doc)
add_paragraph(doc, '3. Masukkan Email dan Password yang valid pada kolom yang disediakan. Anda juga dapat login menggunakan akun Google dengan mengklik tombol "Sign in with Google".')
add_paragraph(doc, '4. Klik tombol Login. Sistem akan memvalidasi data dan mengarahkan pengguna ke halaman Dashboard utama sesuai hak akses perannya.')

add_heading(doc, '2.2 Registrasi dan Lupa Password', level=2)
add_paragraph(doc, '1. Jika belum memiliki akun, klik "Daftar" pada halaman login, isi form yang tersedia, lalu klik "Register".')
add_paragraph(doc, '2. Jika lupa password, klik "Lupa Password?", masukkan email yang terdaftar, dan ikuti instruksi reset password yang dikirim ke email.')
doc.add_page_break()

# --- BAB 3 ---
add_heading(doc, 'BAB 3: PANDUAN OPERASIONAL ADMIN (ROLE 1)', level=1)

add_heading(doc, '3.1 Dashboard dan AI Narasi', level=2)
add_paragraph(doc, '1. Klik menu Dashboard pada bilah navigasi sisi kiri.')
add_paragraph(doc, '2. Di halaman Dashboard, Anda dapat melihat ringkasan data, grafik, dan metrik penting.')
add_screenshot_placeholder(doc)
add_paragraph(doc, '3. Anda dapat menggunakan fitur Generate Narasi AI dengan menekan tombol yang tersedia untuk membuat laporan naratif otomatis berdasarkan data yang difilter.')

add_heading(doc, '3.2 Dashboard Peta', level=2)
add_paragraph(doc, '1. Klik menu Peta Sebaran (atau Dashboard Peta).')
add_paragraph(doc, '2. Halaman ini akan menampilkan peta persebaran data visual sesuai dengan filter wilayah atau indikator yang Anda pilih.')
add_screenshot_placeholder(doc)

add_heading(doc, '3.3 Manajemen Kelola Data', level=2)
add_paragraph(doc, '1. Klik menu Kelola Data.')
add_paragraph(doc, '2. Anda dapat melihat, menambah, mengubah (edit), dan menghapus data Kategori, Subjek, dan Indikator.')
add_paragraph(doc, '3. Untuk mengimpor data secara massal, gunakan fitur Import Excel. Unduh template jika diperlukan, lalu unggah file spreadsheet yang telah diisi.')
add_screenshot_placeholder(doc)

add_heading(doc, '3.4 Manajemen Pengetahuan AI', level=2)
add_paragraph(doc, '1. Klik menu Pengetahuan (Knowledge Base).')
add_paragraph(doc, '2. Di sini Anda dapat mengunggah (upload) dokumen referensi (PDF/Word/Text) yang akan digunakan oleh AI sebagai basis pengetahuan tambahan.')
add_paragraph(doc, '3. Setelah diunggah, klik tombol "Ingest" agar AI mempelajari dokumen tersebut.')
add_screenshot_placeholder(doc)

add_heading(doc, '3.5 Manajemen Pengguna', level=2)
add_paragraph(doc, '1. Klik menu Data Pengguna.')
add_paragraph(doc, '2. Sistem menampilkan daftar akun pengguna beserta role masing-masing.')
add_paragraph(doc, '3. Anda dapat menambah pengguna baru, mengedit role (hak akses), atau menghapus akun pengguna.')
add_screenshot_placeholder(doc)

add_heading(doc, '3.6 Konfigurasi Model AI', level=2)
add_paragraph(doc, '1. Klik menu Model AI.')
add_paragraph(doc, '2. Anda dapat mengecek status koneksi ke model AI (seperti OpenAI atau LLM lokal) dan memperbarui pengaturan API Key atau model yang digunakan.')

doc.add_page_break()

# --- BAB 4 ---
add_heading(doc, 'BAB 4: PANDUAN OPERASIONAL PENANGGUNG JAWAB (ROLE 3)', level=1)
add_paragraph(doc, 'Fungsi operasional Penanggung Jawab sebagian besar sama dengan Admin, namun tanpa hak akses untuk Manajemen Pengguna dan Dashboard Peta (tergantung konfigurasi).')

add_heading(doc, '4.1 Dashboard & Kelola Data', level=2)
add_paragraph(doc, '1. Anda dapat mengakses Dashboard untuk melihat visualisasi data dan melakukan Generate Narasi AI.')
add_paragraph(doc, '2. Melalui menu Kelola Data, Anda memiliki wewenang penuh untuk melakukan CRUD (Create, Read, Update, Delete) pada data indikator, subjek, dan kategori.')

add_heading(doc, '4.2 Mengelola Pengetahuan AI', level=2)
add_paragraph(doc, '1. Anda dapat mengunggah dokumen baru ke basis pengetahuan AI untuk memastikan AI memiliki konteks terbaru saat menghasilkan narasi laporan.')
add_paragraph(doc, '2. Klik "Upload" lalu "Ingest" untuk memasukkan data dokumen ke dalam model.')
add_screenshot_placeholder(doc)

doc.add_page_break()

# --- BAB 5 ---
add_heading(doc, 'BAB 5: PANDUAN OPERASIONAL PENGGUNA UMUM (ROLE 2 & 4)', level=1)
add_paragraph(doc, 'Pengguna Umum hanya memiliki akses untuk melihat informasi dan data tanpa kemampuan mengubah (Read-Only).')

add_heading(doc, '5.1 Melihat Dashboard', level=2)
add_paragraph(doc, '1. Setelah login, Anda akan diarahkan ke halaman Dashboard.')
add_paragraph(doc, '2. Anda dapat melihat grafik statistik dan melakukan penyaringan (filter) data berdasarkan waktu, wilayah, atau kategori indikator.')
add_screenshot_placeholder(doc)

add_heading(doc, '5.2 Melihat Data', level=2)
add_paragraph(doc, '1. Klik menu Lihat Data.')
add_paragraph(doc, '2. Anda dapat melihat tabel data lengkap dan melihat detail (Show Detail) untuk setiap indikator.')
add_paragraph(doc, '3. Tersedia juga fitur "Export" untuk mengunduh data dalam format PDF atau Excel jika dibutuhkan.')
add_screenshot_placeholder(doc)

add_heading(doc, '5.3 Pengaturan Profil', level=2)
add_paragraph(doc, '1. Klik nama Anda atau ikon profil di pojok kanan atas, lalu pilih Pengaturan.')
add_paragraph(doc, '2. Anda dapat memperbarui informasi nama, email, dan mengganti password akun Anda.')
add_screenshot_placeholder(doc)


# Save the document
file_path = 'c:/Herd/pranata/Buku_Panduan_Penggunaan_Pranata.docx'
doc.save(file_path)
print(f"Document saved successfully to {file_path}")
